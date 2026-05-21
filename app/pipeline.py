from __future__ import annotations
import base64
import io
import re
from pathlib import Path

ROLE_PREFIXES = r'(?:арендатор|арендодатель|заказчик|исполнитель|покупатель|продавец|студент|преподаватель|директор|представитель|гражданин|гражданка)'
PERSON_STOPWORDS = {
    'ср', 'средняя', 'маржа', 'минимальная', 'максимальная', 'бонус', 'пакет', 'срок', 'вариант', 'обслуживание',
    'министерство', 'федеральное', 'государственное', 'образовательное', 'учреждение', 'введение', 'условиях',
    'рамках', 'курса', 'алгоритмы', 'структуры', 'данных', 'алгоритм', 'сортировка',
    'российская', 'российской', 'федерация', 'федерации',
}
LOWERCASE_COURSE_WORDS = {'и', 'в', 'на', 'по', 'с', 'со', 'для', 'за', 'из', 'к', 'о', 'об'}

ORG_PATTERNS = [
    (r'Донской государственный технический университет|ДГТУ', 'ДГТУ'),
]
CITY_PATTERNS = [
    (r'Ростов(?:-на-Дону)?', 'Ростов-на-Дону'),
    (r'Москва', 'Москва'),
    (r'Санкт-Петербург', 'Санкт-Петербург'),
]


def extract_text(*, filename: str, content_base64: str) -> tuple[str, str]:
    data = base64.b64decode(content_base64)
    suffix = Path(filename).suffix.lower()
    if suffix in {'.txt', '.md', '.csv'}:
        return data.decode(errors='ignore'), 'text'
    if suffix == '.pdf':
        try:
            from pypdf import PdfReader
            text = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(data)).pages)
            return text, 'pdf'
        except Exception:
            return '', 'ocr_pending'
    if suffix == '.docx':
        from docx import Document
        return '\n'.join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs), 'docx'
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        cells = [str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row if value is not None]
        return '\n'.join(cells), 'xlsx'
    if suffix == '.pptx':
        from pptx import Presentation
        texts = [shape.text for slide in Presentation(io.BytesIO(data)).slides for shape in slide.shapes if hasattr(shape, 'text')]
        return '\n'.join(texts), 'pptx'
    if suffix in {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.webp'}:
        return _ocr_image_bytes(data), 'ocr'
    return data.decode(errors='ignore'), 'binary-fallback'


def _ocr_image_bytes(data: bytes) -> str:
    try:
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)), lang='rus+eng')
    except Exception:
        return ''


def analyze_document_text(text: str) -> dict:
    normalized = ' '.join(text.split())
    facts = _extract_facts(normalized)
    dates = sorted(set(re.findall(r'\b(?:\d{4}-\d{2}-\d{2}|\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\b', normalized)))
    structured_entities = _extract_structured_entities(normalized, facts)
    topic_entities = [] if 'Лабораторная работа' in facts else _extract_topics(normalized)
    summary_items = facts + topic_entities
    summary = '; '.join(dict.fromkeys(summary_items[:8]))
    if not summary:
        sentences = [part.strip() for part in re.split(r'[.!?]', normalized) if part.strip()]
        summary = '. '.join(sentences[:2]) or normalized[:320]
    entities = sorted({entity['name'] for entity in structured_entities} | set(topic_entities))
    events = _extract_events(normalized, dates)
    return {'summary': summary, 'entities': entities, 'structured_entities': structured_entities + [{'kind': 'topic', 'name': topic} for topic in topic_entities], 'dates': dates, 'events': events}


def _extract_facts(text: str) -> list[str]:
    facts: list[str] = []
    if re.search(r'\bдоговор\s+аренды\b', text, re.IGNORECASE):
        facts.append('Договор аренды')
    if re.search(r'\b(?:отчетная|лабораторная)\s+работа\b', text, re.IGNORECASE):
        facts.append('Лабораторная работа')
    for pattern, city in CITY_PATTERNS:
        if re.search(pattern, text, re.IGNORECASE):
            facts.append(city)
    for year in sorted(set(re.findall(r'\b20\d{2}\s*год\b', text, re.IGNORECASE))):
        facts.append(year.replace('  ', ' '))
    m = re.search(r'курс[а]?\s+([А-ЯЁA-Z][^.;:\n]{4,80}?)(?:\s+г\.?\s|\s+20\d{2}|[.;]|$)', text, re.IGNORECASE)
    if m:
        course = _clean_course_title(m.group(1))
        if course:
            facts.append(course)
    for pattern, org in ORG_PATTERNS:
        if re.search(pattern, text, re.IGNORECASE):
            facts.append(org)
    if re.search(r'лабораторн\w+\s+работ', text, re.IGNORECASE):
        facts.append('Лабораторные работы')
    return facts


def _extract_structured_entities(text: str, facts: list[str]) -> list[dict[str, str]]:
    entities: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    is_lab_report = 'Лабораторная работа' in facts
    def add(kind: str, name: str) -> None:
        clean = _clean_name(name)
        if kind == 'person':
            clean = _normalize_person(clean)
            if not clean:
                return
        key = (kind, clean.lower())
        if key not in seen:
            seen.add(key)
            entities.append({'kind': kind, 'name': clean})
    person_pattern = rf'\b(?:(?i:{ROLE_PREFIXES})\s+)?(?:[А-ЯЁ][а-яё]{{2,}}\s+){{1,2}}[А-ЯЁ][а-яё]{{2,}}\b'
    for match in re.finditer(person_pattern, text):
        add('person', match.group(0))
    for match in re.finditer(r'\b(?:ООО|АО|ПАО|ЗАО|ИП|ОАО)\s+[«"]?[^,.;:\n]{2,80}', text):
        add('company', match.group(0))
    for pattern, org in ORG_PATTERNS:
        if re.search(pattern, text, re.IGNORECASE):
            add('company', org)
    for pattern, city in CITY_PATTERNS:
        if re.search(pattern, text, re.IGNORECASE):
            add('city', city)
    if not is_lab_report:
        for match in re.finditer(r'\b(?:проект|Project)\s+[«"]?[^,.;:\n]{2,70}', text, re.IGNORECASE):
            add('project', match.group(0))
        for match in re.finditer(r'\b\d[\d\s.,]{1,18}\s*(?:₽|руб\.?|евро|eur|€|usd|\$)\b', text, re.IGNORECASE):
            amount = match.group(0)
            add('finance', amount)
            bare = re.match(r'\d[\d\s.,]{0,18}', amount)
            if bare:
                add('finance', bare.group(0).strip())
        for match in re.finditer(r'\b(?:сумма|залог|депозит|штраф|оплата|сч[её]т(?!чик)|налог)[^,.;:\n]{0,60}', text, re.IGNORECASE):
            add('finance', match.group(0))
    return entities


def _extract_topics(text: str) -> list[str]:
    topics = []
    rules = {'Договор': r'\bдоговор\b', 'Недвижимость': r'\b(?:квартир|аренд|недвижим|помещени)\w*', 'Медицина': r'\b(?:медицин|клиник|врач|при[её]м\s+врача|медицинск\w+\s+при[её]м)\w*', 'Налоги': r'\b(?:налог|деклараци|фнс)\w*', 'Страхование': r'\b(?:страхов|полис)\w*', 'Финансы': r'\b(?:сч[её]т(?!чик)|оплат|банк|выписк|плат[её]ж)\w*'}
    for name, pattern in rules.items():
        if re.search(pattern, text, re.IGNORECASE):
            topics.append(name)
    return topics


def _extract_events(text: str, dates: list[str]) -> list[dict[str, str | None]]:
    events = []
    for date in dates:
        idx = text.find(date)
        context = text[max(0, idx - 90): idx + 140] if idx >= 0 else text[:220]
        title = 'Дата из документа'
        if re.search(r'оплат|сч[её]т|плат[её]ж', context, re.IGNORECASE):
            title = 'Срок оплаты'
        elif re.search(r'окончан|истека|продлен|страхов', context, re.IGNORECASE):
            title = 'Срок действия документа'
        elif re.search(r'встреч|при[её]м|визит', context, re.IGNORECASE):
            title = 'Встреча из документа'
        events.append({'title': title, 'starts_at': date, 'description': _clean_name(context) or None})
    return events


def _normalize_person(value: str) -> str | None:
    value = re.sub(rf'^(?:{ROLE_PREFIXES})\s+', '', value.strip(), flags=re.IGNORECASE)
    parts = [p for p in value.split() if p]
    if len(parts) < 2 or len(parts) > 3:
        return None
    if any(len(p) < 3 for p in parts):
        return None
    lowered = {p.lower() for p in parts}
    if lowered & PERSON_STOPWORDS:
        return None
    # Reject phrases with adjective-like middle/business terms instead of plausible names.
    if any(not re.fullmatch(r'[А-ЯЁ][а-яё]{2,}', p) for p in parts):
        return None
    if any(p.lower().endswith(('ая', 'ое', 'ый', 'ий')) for p in parts[1:]) and len(parts) < 3:
        return None
    return ' '.join(parts)


def _clean_course_title(value: str) -> str | None:
    value = re.sub(r'\bг\.?$', '', _clean_name(value), flags=re.IGNORECASE).strip()
    words = [word for word in re.split(r'\s+', value) if word]
    words = [word for word in words if word.lower() not in {'г', 'год'}]
    if not words:
        return None
    result = []
    for index, word in enumerate(words):
        low = word.lower()
        result.append(word[:1].upper() + word[1:].lower() if index == 0 else low)
    title = ' '.join(result)
    return title.strip(' .;:,') or None


def _title_case(value: str) -> str:
    return ' '.join(part[:1].upper() + part[1:].lower() if part.isupper() else part for part in value.split())


def _clean_name(value: str) -> str:
    return re.sub(r'\s+', ' ', value.strip(' «"”.,;:()[]')).strip()
